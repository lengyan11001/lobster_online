"""文本渲染器 - 含 CJK 字体 workaround"""

from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from lxml import etree

from ppt_maker.models import TextElement, TextRun, TextStyle, LayoutZone, ThemeModel


class TextRenderer:
    """文本渲染器"""

    def __init__(self, theme: ThemeModel):
        self.theme = theme

    def render_text(self, shape, element: TextElement, zone: LayoutZone):
        """渲染文本元素到形状"""
        tf = shape.text_frame
        tf.word_wrap = True
        tf.auto_size = None

        # 设置边距
        tf.margin_left = Emu(int(zone.width * 0.03))
        tf.margin_right = Emu(int(zone.width * 0.03))
        tf.margin_top = Emu(int(zone.height * 0.02))
        tf.margin_bottom = Emu(int(zone.height * 0.02))

        style = element.style or TextStyle()

        # 收集所有 runs
        all_runs = []
        if element.text:
            all_runs.append(TextRun(text=element.text, style=None))
        all_runs.extend(element.runs)

        if not all_runs:
            return

        first_paragraph = True
        for run_data in all_runs:
            if first_paragraph:
                p = tf.paragraphs[0]
                first_paragraph = False
            else:
                p = tf.add_paragraph()

            run_style = run_data.style or style
            p.alignment = self._alignment_value(run_style.alignment)

            # 设置行间距
            p.line_spacing = run_style.line_spacing

            # 设置段落间距
            p.space_after = Pt(4)

            # 设置列表缩进
            if run_style.bullet:
                p.level = run_style.bullet_level

            run = p.add_run()
            run.text = run_data.text

            # 应用字体样式 (含 CJK workaround)
            self._apply_font(run, run_style)

    def _apply_font(self, run, style: TextStyle):
        """应用字体样式，包括 CJK 东亚字体 workaround"""
        run.font.size = Pt(style.font_size)
        run.font.bold = style.bold
        run.font.italic = style.italic

        # 颜色
        color_str = style.color.lstrip("#")
        if len(color_str) == 6:
            run.font.color.rgb = RGBColor.from_string(color_str)

        # === CJK 字体 workaround ===
        # python-pptx 的 run.font.name 只设置 a:latin (西文字体)
        # 需要额外设置 a:ea 元素来指定东亚字体
        run.font.name = style.font_name_latin

        rPr = run._r.get_or_add_rPr()

        # 设置东亚字体 a:ea
        ea = rPr.find(qn("a:ea"))
        if ea is None:
            ea = etree.SubElement(rPr, qn("a:ea"))
        ea.set("typeface", style.font_name)

        # 确保西文字体 a:latin 一致
        latin = rPr.find(qn("a:latin"))
        if latin is None:
            latin = etree.SubElement(rPr, qn("a:latin"))
        latin.set("typeface", style.font_name_latin)

    @staticmethod
    def _alignment_value(alignment: str):
        """将字符串对齐方式转为 python-pptx 枚举"""
        from pptx.enum.text import PP_ALIGN

        return {
            "left": PP_ALIGN.LEFT,
            "center": PP_ALIGN.CENTER,
            "right": PP_ALIGN.RIGHT,
            "justify": PP_ALIGN.JUSTIFY,
        }.get(alignment, PP_ALIGN.LEFT)
